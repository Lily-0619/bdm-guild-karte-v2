# -*- coding: utf-8 -*-
"""Per-guild comment source data: combine guild aggregates with member detail.

This is the first step toward AI-written guild-card comments.  It reads the
session summary (``analysis/summary_<date>.xlsx``) -- which already holds each
guild's aggregate metrics *and* whole-server comparisons -- and joins it with
the individual member list resolved through the summary's ``source_file`` column
(robust against night-crossing collection).  For each guild it writes:

* one Excel sheet (human viewable), and
* one Markdown file (ready to feed an AI later),

merging "guild info" and "individual info" that were previously separate.

It intentionally reuses the Japanese phrasing helpers already in ``analyze.py``
(``describe_power_features`` / ``describe_growth`` / ``rank_position_label``) so
the comment material matches the existing rule-based ``auto_comment`` wording.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Any, Iterable, Optional

from openpyxl import Workbook, load_workbook

try:
    from .paths import ANALYSIS_DIR, PROJECT_ROOT, ensure_dirs
    from .analyze import (
        classify_power_type,
        describe_growth,
        describe_power_features,
        median_numeric,
        rank_position_label,
        to_number,
    )
except ImportError:  # 直接実行された場合のため
    from paths import ANALYSIS_DIR, PROJECT_ROOT, ensure_dirs  # type: ignore
    from analyze import (  # type: ignore
        classify_power_type,
        describe_growth,
        describe_power_features,
        median_numeric,
        rank_position_label,
        to_number,
    )


SUMMARY_DATE_RE = re.compile(r"summary_(\d{4})-(\d{2})-(\d{2})\.xlsx$", re.IGNORECASE)

# 個人一覧で見せる主要数値（key, 表示ラベル）。
MEMBER_DISPLAY_COLUMNS = (
    ("rank", "順位"),
    ("player_name", "家門名"),
    ("level", "レベル"),
    ("cpm", "CPM"),
    ("growth", "前回比"),
    ("fcp", "FCP"),
)

# Excel上段で見せるギルド集計＋全体比較（key, 表示ラベル）。
GUILD_SUMMARY_FIELDS = (
    ("rank_by_avg_cpm", "平均CPM順位"),
    ("rank_position", "全体での位置"),
    ("member_count", "人数"),
    ("avg_cpm", "平均CPM"),
    ("median_cpm", "中央値CPM"),
    ("avg_cpm_diff_from_all_avg", "全体平均との差"),
    ("avg_cpm_growth", "前回比(平均CPM)"),
    ("avg_cpm_growth_rate", "前回比(率)"),
    ("growth_rank", "伸びランク"),
    ("growth_vs_all_avg", "全体平均比の伸び"),
    ("power_type", "戦力タイプ"),
    ("top10_avg_cpm", "上位10平均"),
    ("max_cpm", "最大CPM"),
    ("min_cpm", "最小CPM"),
    ("stdev_cpm", "標準偏差"),
    ("all_time_win_rate", "通算勝率"),
    ("node_win_rate", "ノード勝率"),
    ("siege_win_rate", "シージ勝率"),
)


def safe_filename(value: str) -> str:
    """Windows/Macで使えない文字を避けたファイル名にする。"""
    sanitized = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(value).strip())
    sanitized = re.sub(r"\s+", "_", sanitized).strip(" ._")
    return sanitized or "unknown_guild"


def summary_date_from_path(path: Path) -> str:
    match = SUMMARY_DATE_RE.search(path.name)
    if not match:
        return ""
    return f"{match.group(1)}-{match.group(2)}-{match.group(3)}"


def list_summary_paths() -> list[Path]:
    """Return summary workbooks sorted oldest -> newest by date in filename."""
    if not ANALYSIS_DIR.exists():
        return []
    dated = [(summary_date_from_path(p), p) for p in ANALYSIS_DIR.glob("summary_*.xlsx")]
    dated = [(date_str, p) for date_str, p in dated if date_str]
    return [p for _, p in sorted(dated, key=lambda item: item[0])]


def summary_path_for_date(date_str: str) -> Optional[Path]:
    for path in list_summary_paths():
        if summary_date_from_path(path) == date_str:
            return path
    return None


def normalize_header(value: Any) -> str:
    return re.sub(r"[\s_\-　（）()]+", "", str(value).strip().lower()) if value is not None else ""


def rows_to_dicts(sheet) -> list[dict[str, Any]]:
    """Turn a worksheet (header row + data rows) into a list of dicts."""
    rows = sheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        return []
    keys = [str(cell).strip() if cell is not None else f"col{i}" for i, cell in enumerate(header)]
    records: list[dict[str, Any]] = []
    for row in rows:
        if all(cell in (None, "") for cell in row):
            continue
        records.append({keys[i]: row[i] if i < len(row) else None for i in range(len(keys))})
    return records


def load_guild_metrics(summary_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(summary_path, data_only=True, read_only=True)
    try:
        if "guild_metrics" not in workbook.sheetnames:
            return []
        return rows_to_dicts(workbook["guild_metrics"])
    finally:
        workbook.close()


def resolve_source_file(row: dict[str, Any]) -> Optional[Path]:
    source_file = row.get("source_file")
    if not source_file:
        return None
    path = Path(str(source_file))
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    return path if path.exists() else None


MEMBER_COLUMN_ALIASES = {
    "rank": ("rank", "順位", "ランク"),
    "player_name": ("player_name", "family_name", "name", "家門名", "プレイヤー", "名前"),
    "level": ("level", "レベル", "lv"),
    "cpm": ("cpm", "戦闘力"),
    "fcp": ("fcp", "family_cp", "家門戦闘力"),
}


def detect_member_columns(header: Iterable[Any]) -> dict[str, int]:
    normalized = [normalize_header(cell) for cell in header]
    mapping: dict[str, int] = {}
    for field, aliases in MEMBER_COLUMN_ALIASES.items():
        for index, head in enumerate(normalized):
            if head and any(normalize_header(alias) in head for alias in aliases):
                mapping[field] = index
                break
    # 固定順フォールバック: rank, player_name, level, cpm, fcp
    defaults = {"rank": 0, "player_name": 1, "level": 2, "cpm": 3, "fcp": 4}
    for field, index in defaults.items():
        mapping.setdefault(field, index)
    return mapping


def load_members(workbook_path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    try:
        if "members" not in workbook.sheetnames:
            return []
        sheet = workbook["members"]
        iterator = sheet.iter_rows(values_only=True)
        header = next(iterator, None)
        if header is None:
            return []
        columns = detect_member_columns(header)
        members: list[dict[str, Any]] = []
        for row in iterator:
            name = row[columns["player_name"]] if columns["player_name"] < len(row) else None
            if name in (None, ""):
                continue
            members.append(
                {
                    "rank": _cell(row, columns["rank"]),
                    "player_name": str(name).strip(),
                    "level": _cell(row, columns["level"]),
                    "cpm": to_number(_cell(row, columns["cpm"])),
                    "fcp": to_number(_cell(row, columns["fcp"])),
                }
            )
        return members
    finally:
        workbook.close()


def _cell(row: tuple, index: int) -> Any:
    return row[index] if index is not None and index < len(row) else None


def build_old_cpm_index(old_summary_path: Optional[Path]) -> dict[tuple[str, str], float]:
    """Map (guild_name, player_name) -> previous CPM for growth calculation."""
    if old_summary_path is None:
        return {}
    index: dict[tuple[str, str], float] = {}
    for row in load_guild_metrics(old_summary_path):
        guild_name = str(row.get("guild_name") or "").strip()
        source = resolve_source_file(row)
        if not guild_name or source is None:
            continue
        for member in load_members(source):
            cpm = member.get("cpm")
            if cpm is not None:
                index[(guild_name, member["player_name"])] = cpm
    return index


def build_guild_record(
    row: dict[str, Any],
    *,
    rank_by_avg_cpm: int,
    stdev_median: float,
    old_cpm_index: dict[tuple[str, str], float],
) -> dict[str, Any]:
    """Combine one guild's aggregate metrics with its member detail."""
    guild_name = str(row.get("guild_name") or "").strip()
    source = resolve_source_file(row)
    members = load_members(source) if source else []

    enriched_members: list[dict[str, Any]] = []
    for member in members:
        previous = old_cpm_index.get((guild_name, member["player_name"]))
        current = member.get("cpm")
        growth = current - previous if (previous is not None and current is not None) else None
        enriched_members.append({**member, "growth": growth})
    enriched_members.sort(key=lambda m: (m.get("cpm") is None, -(m.get("cpm") or 0)))

    if not row.get("power_type"):
        row = {**row, "power_type": classify_power_type(row, stdev_median)}

    grown = [m for m in enriched_members if m.get("growth") is not None]
    growth_top = sorted(grown, key=lambda m: m["growth"], reverse=True)[:5]
    growth_bottom = [m for m in sorted(grown, key=lambda m: m["growth"]) if m["growth"] <= 0][:5]

    return {
        "guild_name": guild_name,
        "summary_row": row,
        "rank_by_avg_cpm": rank_by_avg_cpm,
        "rank_position": rank_position_label(row.get("rank_percentile")),
        "members": enriched_members,
        "growth_top": growth_top,
        "growth_bottom": growth_bottom,
        "phrases": {
            "position": _position_phrase(row, rank_by_avg_cpm),
            "power": describe_power_features(row, stdev_median),
            "growth": describe_growth(row),
        },
        "missing_source": source is None,
    }


def _position_phrase(row: dict[str, Any], rank_by_avg_cpm: int) -> str:
    position = rank_position_label(row.get("rank_percentile"))
    return f"平均CPMは全体{rank_by_avg_cpm}位で、{position}に位置します。"


def build_guild_records(
    new_summary_path: Path, old_summary_path: Optional[Path]
) -> list[dict[str, Any]]:
    metrics = load_guild_metrics(new_summary_path)
    ranked = sorted(
        metrics,
        key=lambda r: to_number(r.get("avg_cpm")) if to_number(r.get("avg_cpm")) is not None else -1,
        reverse=True,
    )
    rank_by_guild = {id(row): position for position, row in enumerate(ranked, start=1)}
    stdev_median = median_numeric(row.get("stdev_cpm") for row in metrics)
    old_cpm_index = build_old_cpm_index(old_summary_path)

    records = [
        build_guild_record(
            row,
            rank_by_avg_cpm=rank_by_guild[id(row)],
            stdev_median=stdev_median,
            old_cpm_index=old_cpm_index,
        )
        for row in metrics
        if str(row.get("guild_name") or "").strip()
    ]
    records.sort(key=lambda rec: rec["rank_by_avg_cpm"])
    return records


# ---- 出力: Excel ----------------------------------------------------------

def _format_value(value: Any) -> Any:
    number = to_number(value)
    if number is not None and value not in (None, ""):
        return round(number, 2) if isinstance(number, float) and not number.is_integer() else number
    return "" if value is None else value


def write_excel(records: list[dict[str, Any]], out_path: Path) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)
    used_titles: set[str] = set()
    for record in records:
        title = _unique_sheet_title(record["guild_name"], used_titles)
        sheet = workbook.create_sheet(title=title)
        row_index = 1
        sheet.cell(row=row_index, column=1, value="ギルド").font = _bold()
        sheet.cell(row=row_index, column=2, value=record["guild_name"])
        row_index += 1

        summary_row = record["summary_row"]
        for key, label in GUILD_SUMMARY_FIELDS:
            value = record[key] if key in record else summary_row.get(key)
            sheet.cell(row=row_index, column=1, value=label)
            sheet.cell(row=row_index, column=2, value=_format_value(value))
            row_index += 1

        for phrase_label, phrase_key in (("位置", "position"), ("構成傾向", "power"), ("伸び", "growth")):
            sheet.cell(row=row_index, column=1, value=phrase_label)
            sheet.cell(row=row_index, column=2, value=record["phrases"][phrase_key])
            row_index += 1

        row_index += 1  # 空行
        header_row = row_index
        for col, (_key, label) in enumerate(MEMBER_DISPLAY_COLUMNS, start=1):
            sheet.cell(row=header_row, column=col, value=label).font = _bold()
        row_index += 1
        for member in record["members"]:
            for col, (key, _label) in enumerate(MEMBER_DISPLAY_COLUMNS, start=1):
                sheet.cell(row=row_index, column=col, value=_format_value(member.get(key)))
            row_index += 1

        sheet.column_dimensions["A"].width = 18
        sheet.column_dimensions["B"].width = 22

    if not workbook.sheetnames:
        workbook.create_sheet(title="empty")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)


def _unique_sheet_title(guild_name: str, used: set[str]) -> str:
    base = re.sub(r'[\[\]\:\*\?/\\]', "_", guild_name).strip() or "guild"
    title = base[:31]
    counter = 1
    while title in used or not title:
        suffix = f"_{counter}"
        title = base[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(title)
    return title


def _bold():
    from openpyxl.styles import Font

    return Font(bold=True)


# ---- 出力: Markdown -------------------------------------------------------

def _fmt(value: Any) -> str:
    number = to_number(value)
    if number is not None and value not in (None, ""):
        if isinstance(number, float) and not number.is_integer():
            return f"{number:,.2f}"
        return f"{int(number):,}"
    return "不明" if value in (None, "") else str(value)


def build_guild_markdown(record: dict[str, Any], new_date: str, old_date: Optional[str]) -> str:
    row = record["summary_row"]
    lines: list[str] = []
    lines.append(f"# {record['guild_name']} コメント元データ")
    lines.append("")
    period = f"集計日: {new_date}"
    if old_date:
        period += f" ／ 前回比較日: {old_date}"
    lines.append(period)
    lines.append("")

    lines.append("## 全体の中での位置づけ（AI向け下書き）")
    lines.append(f"- {record['phrases']['position']}")
    lines.append(f"- 構成傾向: {record['phrases']['power']}")
    lines.append(f"- 伸び: {record['phrases']['growth']}")
    lines.append("")

    lines.append("## ギルド集計 ＋ 全体比較")
    for key, label in GUILD_SUMMARY_FIELDS:
        value = record[key] if key in record else row.get(key)
        if key in ("rank_position", "power_type"):
            lines.append(f"- {label}: {value if value not in (None, '') else '不明'}")
        else:
            lines.append(f"- {label}: {_fmt(value)}")
    lines.append("")

    if record["growth_top"]:
        lines.append("## 個人ハイライト・伸び上位")
        for member in record["growth_top"]:
            lines.append(f"- {member['player_name']}: 前回比 +{_fmt(member['growth'])}（CPM {_fmt(member['cpm'])}）")
        lines.append("")
    if record["growth_bottom"]:
        lines.append("## 個人ハイライト・停滞/減少")
        for member in record["growth_bottom"]:
            lines.append(f"- {member['player_name']}: 前回比 {_fmt(member['growth'])}（CPM {_fmt(member['cpm'])}）")
        lines.append("")

    top_members = record["members"][:10]
    if top_members:
        lines.append("## 上位CPMメンバー")
        for member in top_members:
            growth = member.get("growth")
            growth_text = f" / 前回比 {_fmt(growth)}" if growth is not None else ""
            lines.append(f"- {member['player_name']}: CPM {_fmt(member['cpm'])}{growth_text}")
        lines.append("")

    lines.append(f"## 全メンバー一覧（{len(record['members'])}人）")
    lines.append("| 順位 | 家門名 | レベル | CPM | 前回比 | FCP |")
    lines.append("| --- | --- | --- | --- | --- | --- |")
    for member in record["members"]:
        growth = member.get("growth")
        growth_text = _fmt(growth) if growth is not None else "-"
        lines.append(
            f"| {_fmt(member.get('rank'))} | {member['player_name']} | {_fmt(member.get('level'))} "
            f"| {_fmt(member.get('cpm'))} | {growth_text} | {_fmt(member.get('fcp'))} |"
        )
    lines.append("")

    lines.append("## AIへの注記")
    lines.append(
        "上のデータを材料に、このギルド向けの自然な日本語コメントを書いてください。"
        "数字を全部使う必要はありません。全体と比べた強さ・伸び・構成の雰囲気（まったり/本気度）・"
        "注目メンバーが伝わる、事務的すぎず煽りすぎない文章にしてください。"
    )
    lines.append("")
    return "\n".join(lines)


def write_markdown(records: list[dict[str, Any]], out_dir: Path, new_date: str, old_date: Optional[str]) -> None:
    out_dir.mkdir(parents=True, exist_ok=True)
    index_lines = [f"# コメント元データ {new_date}", ""]
    if old_date:
        index_lines.append(f"前回比較日: {old_date}")
        index_lines.append("")
    for record in records:
        file_name = f"{safe_filename(record['guild_name'])}.md"
        (out_dir / file_name).write_text(
            build_guild_markdown(record, new_date, old_date), encoding="utf-8"
        )
        index_lines.append(
            f"- [{record['guild_name']}]({file_name}) "
            f"（{record['rank_by_avg_cpm']}位 / {record['rank_position']} / {record['summary_row'].get('power_type', '')}）"
        )
    index_lines.append("")
    (out_dir / "index.md").write_text("\n".join(index_lines), encoding="utf-8")


# ---- 実行 -----------------------------------------------------------------

def run(new_date: Optional[str] = None, old_date: Optional[str] = None) -> dict[str, Path]:
    ensure_dirs()
    summaries = list_summary_paths()
    if not summaries:
        raise FileNotFoundError(
            f"analysis/summary_*.xlsx が見つかりません。先に「一覧作成」を実行してください: {ANALYSIS_DIR}"
        )

    if new_date:
        new_summary = summary_path_for_date(new_date)
        if new_summary is None:
            raise FileNotFoundError(f"指定日のサマリーが見つかりません: summary_{new_date}.xlsx")
    else:
        new_summary = summaries[-1]
    new_date = summary_date_from_path(new_summary)

    if old_date:
        old_summary = summary_path_for_date(old_date)
    else:
        earlier = [p for p in summaries if summary_date_from_path(p) < new_date]
        old_summary = earlier[-1] if earlier else None
    old_date_str = summary_date_from_path(old_summary) if old_summary else None

    records = build_guild_records(new_summary, old_summary)

    excel_path = ANALYSIS_DIR / f"comment_source_{new_date}.xlsx"
    markdown_dir = ANALYSIS_DIR / "comment_source" / new_date
    write_excel(records, excel_path)
    write_markdown(records, markdown_dir, new_date, old_date_str)

    print("コメント元データを作成しました。")
    print(f"対象サマリー: {new_summary.name}" + (f" / 前回: {old_summary.name}" if old_summary else " / 前回: なし"))
    print(f"ギルド数: {len(records)}")
    print(f"Excel: {excel_path}")
    print(f"Markdown: {markdown_dir}")
    return {"excel": excel_path, "markdown_dir": markdown_dir}


def main(argv: Optional[Iterable[str]] = None) -> int:
    parser = argparse.ArgumentParser(
        description="サマリーからギルド単位のコメント元データ（Excel＋Markdown）を作成します。"
    )
    parser.add_argument("--new-date", default=None, help="対象サマリー日付 (YYYY-MM-DD)。省略時は最新。")
    parser.add_argument("--old-date", default=None, help="前回比較サマリー日付 (YYYY-MM-DD)。省略時は直近過去。")
    args = parser.parse_args(list(argv) if argv is not None else None)
    run(new_date=args.new_date, old_date=args.old_date)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
