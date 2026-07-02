# -*- coding: utf-8 -*-
"""Render a filled karte workbook range to PNG with Pillow (no Excel needed).

``make_card.py`` は従来 Windows + デスクトップExcel の COM 経由でカルテPNGを
出力していた。このモジュールは openpyxl でセルの値・結合・塗り・罫線・
フォント・配置を読み取り、Pillow で直接描画することで Excel 依存を無くす。

再現しているもの:
    列幅/行高（Excel単位→ピクセル換算）、結合セル、単色塗り（テーマ色＋tint対応）、
    罫線（hair/thin/medium/thick）、フォント（名前解決・太字・色）、
    配置（左右中央/上下・折り返し・隣接空セルへのはみ出し）、数値書式（#,##0 / 0 / %）。

フォントは Windows のフォントフォルダから名前で解決する。テンプレートが使う
フォントが見つからない場合は日本語フォールバック（游明朝/游ゴシック/メイリオ）へ
落ちる。特定のフォントを固定したい場合は ``config/karte_fonts.json`` に
``{"フォント名": "C:/path/to/font.ttf"}`` を書けば最優先で使われる。
"""

from __future__ import annotations

import json
import logging
import os
import re
import unicodedata
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import range_boundaries
from PIL import Image, ImageDraw, ImageFont

try:
    from .paths import CONFIG_DIR
except ImportError:  # 直接実行された場合のため
    from paths import CONFIG_DIR  # type: ignore

logger = logging.getLogger(__name__)

FONT_CONFIG_PATH = CONFIG_DIR / "karte_fonts.json"

# Excel の列幅単位（標準フォント基準の文字数）→ピクセル換算係数。
CHAR_PIXELS = 7
COL_PADDING_PIXELS = 5
DEFAULT_COL_WIDTH = 8.43
DEFAULT_ROW_HEIGHT_PT = 15.0
POINTS_TO_PIXELS = 96 / 72

# theme="n" のインデックス→clrScheme要素。Excel は dk1/lt1, dk2/lt2 を入れ替えて扱う。
THEME_SLOT_ORDER = (
    "lt1", "dk1", "lt2", "dk2",
    "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
    "hlink", "folHlink",
)

BORDER_WIDTHS = {
    "hair": 1, "thin": 1, "dotted": 1, "dashed": 1,
    "medium": 2, "double": 2, "mediumDashed": 2,
    "thick": 3,
}

# 日本語グリフを含まないことが分かっているフォント。CJK文字を含むテキストに
# これらが指定されていたら日本語フォールバックへ差し替える（豆腐化防止）。
LATIN_ONLY_FAMILIES = {"dancing script", "dancingscript"}

# テンプレートの日本語フォント名 → 英語ファミリー名の別名。
FONT_ALIASES = {
    "ms p明朝": ["ms pmincho"],
    "ms 明朝": ["ms mincho"],
    "ms pゴシック": ["ms pgothic"],
    "ms ゴシック": ["ms gothic"],
    "hgp行書体": ["hgpgyoshotai", "hgp行書体"],
    "hg行書体": ["hggyoshotai"],
    "メイリオ": ["meiryo"],
    "游ゴシック": ["yu gothic"],
    "游明朝": ["yu mincho"],
    "かなたとひなた漢": ["kanata to hinata kan", "かなたとひなた"],
}

FALLBACK_FAMILIES = [
    "yu mincho", "yu gothic ui", "yu gothic", "meiryo",
    "ms pgothic", "ms gothic", "ms pmincho", "ms mincho",
]

_CJK_RE = re.compile(r"[　-ヿ㐀-鿿豈-﫿＀-￯]")


def _norm_family(name: str) -> str:
    """フォント名を比較用に正規化（全角→半角、小文字化、空白圧縮）。"""
    text = unicodedata.normalize("NFKC", str(name or "")).strip().lower()
    return re.sub(r"\s+", " ", text)


# ---------------------------------------------------------------------------
# 色の解決
# ---------------------------------------------------------------------------

def load_theme_colors(xlsx_path: Path) -> dict[str, str]:
    """xl/theme/theme1.xml から clrScheme の色（RRGGBB）を読む。"""
    colors: dict[str, str] = {}
    try:
        with zipfile.ZipFile(xlsx_path) as archive:
            xml_text = archive.read("xl/theme/theme1.xml").decode("utf-8", "replace")
    except (OSError, KeyError, zipfile.BadZipFile):
        return colors
    for slot in {"dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3",
                 "accent4", "accent5", "accent6", "hlink", "folHlink"}:
        match = re.search(
            rf"<a:{slot}>.*?(?:val=\"([0-9A-Fa-f]{{6}})\"|lastClr=\"([0-9A-Fa-f]{{6}})\")",
            xml_text,
            re.DOTALL,
        )
        if match:
            colors[slot] = (match.group(1) or match.group(2)).upper()
    return colors


def apply_tint(rgb: tuple[int, int, int], tint: float) -> tuple[int, int, int]:
    """Excel の tint をRGB近似で適用する（正=白寄せ、負=黒寄せ）。"""
    if not tint:
        return rgb
    if tint > 0:
        return tuple(int(round(c + (255 - c) * tint)) for c in rgb)  # type: ignore[return-value]
    return tuple(int(round(c * (1 + tint))) for c in rgb)  # type: ignore[return-value]


def _hex_to_rgb(hex6: str) -> tuple[int, int, int]:
    return tuple(int(hex6[i : i + 2], 16) for i in (0, 2, 4))  # type: ignore[return-value]


def resolve_color(color: Any, theme_colors: dict[str, str]) -> tuple[int, int, int] | None:
    """openpyxl の Color オブジェクトを (r, g, b) にする。解決不能なら None。"""
    if color is None:
        return None
    try:
        color_type = color.type
    except AttributeError:
        return None
    if color_type == "rgb":
        rgb = color.rgb
        if not isinstance(rgb, str) or len(rgb) < 6:
            return None
        return _hex_to_rgb(rgb[-6:])
    if color_type == "theme":
        theme_index = color.theme
        if not isinstance(theme_index, int) or not (0 <= theme_index < len(THEME_SLOT_ORDER)):
            return None
        slot = THEME_SLOT_ORDER[theme_index]
        base_hex = theme_colors.get(slot)
        if base_hex is None:
            return None
        return apply_tint(_hex_to_rgb(base_hex), float(color.tint or 0.0))
    if color_type == "indexed":
        # 旧式インデックス色。テンプレートでは実質使われないため主要色のみ対応。
        legacy = {0: "000000", 1: "FFFFFF", 2: "FF0000", 3: "00FF00", 4: "0000FF",
                  8: "000000", 9: "FFFFFF", 64: None, 65: None}
        hex6 = legacy.get(color.indexed, "000000")
        return _hex_to_rgb(hex6) if hex6 else None
    return None


# ---------------------------------------------------------------------------
# フォントの解決
# ---------------------------------------------------------------------------

_font_index_cache: dict[str, dict[str, tuple[str, int]]] | None = None
_font_object_cache: dict[tuple[str, int, int], ImageFont.FreeTypeFont] = {}


def _font_directories() -> list[Path]:
    dirs = [Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"]
    local = os.environ.get("LOCALAPPDATA")
    if local:
        dirs.append(Path(local) / "Microsoft" / "Windows" / "Fonts")
    return [d for d in dirs if d.is_dir()]


def _scan_installed_fonts() -> dict[str, dict[str, tuple[str, int]]]:
    """インストール済みフォントを {family: {style: (path, index)}} に索引化する。"""
    index: dict[str, dict[str, tuple[str, int]]] = {}
    for directory in _font_directories():
        for path in directory.iterdir():
            suffix = path.suffix.lower()
            if suffix not in (".ttf", ".otf", ".ttc"):
                continue
            max_faces = 8 if suffix == ".ttc" else 1
            for face_index in range(max_faces):
                try:
                    font = ImageFont.truetype(str(path), size=12, index=face_index)
                    family, style = font.getname()
                except OSError:
                    break
                except Exception:  # noqa: BLE001 - 壊れたフォントは無視。
                    break
                family_key = _norm_family(family)
                style_key = _norm_family(style)
                index.setdefault(family_key, {}).setdefault(
                    style_key, (str(path), face_index)
                )
    return index


def _installed_fonts() -> dict[str, dict[str, tuple[str, int]]]:
    global _font_index_cache
    if _font_index_cache is None:
        _font_index_cache = _scan_installed_fonts()
        logger.info("フォント索引を作成しました: %d ファミリー", len(_font_index_cache))
    return _font_index_cache


def _load_font_config() -> dict[str, str]:
    try:
        raw = json.loads(FONT_CONFIG_PATH.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {}
    return {_norm_family(k): str(v) for k, v in raw.items() if v}


_missing_font_warned: set[str] = set()


def resolve_font_file(family: str, bold: bool) -> tuple[str, int] | None:
    """フォント名から (ファイルパス, faceインデックス) を解決する。"""
    family_key = _norm_family(family)

    config = _load_font_config()
    override = config.get(family_key)
    if override and Path(override).is_file():
        return override, 0

    installed = _installed_fonts()
    candidates = [family_key] + FONT_ALIASES.get(family_key, [])
    for candidate in candidates:
        styles = installed.get(_norm_family(candidate))
        if not styles:
            continue
        if bold:
            for style_name, entry in styles.items():
                if "bold" in style_name and "italic" not in style_name:
                    return entry
        for preferred in ("regular", "normal", "標準", "book", "medium"):
            if preferred in styles:
                return styles[preferred]
        return next(iter(styles.values()))

    if family_key not in _missing_font_warned:
        _missing_font_warned.add(family_key)
        logger.warning(
            "フォント '%s' が見つからないため代替を使います"
            "（config/karte_fonts.json で指定できます）。",
            family,
        )
    for fallback in FALLBACK_FAMILIES:
        styles = installed.get(fallback)
        if styles:
            return next(iter(styles.values()))
    return None


def get_font(family: str, size_px: int, bold: bool, text: str = "") -> tuple[Any, bool]:
    """描画用フォントを返す。戻り値は (font, fake_boldが必要か)。"""
    # 日本語を含むのにラテン専用フォント指定なら、日本語フォールバックへ。
    if text and _CJK_RE.search(text) and _norm_family(family) in LATIN_ONLY_FAMILIES:
        for fallback in FALLBACK_FAMILIES:
            if _installed_fonts().get(fallback):
                family = fallback
                break

    entry = resolve_font_file(family, bold)
    if entry is None:
        return ImageFont.load_default(), False
    path, face_index = entry
    cache_key = (path, face_index, size_px)
    font = _font_object_cache.get(cache_key)
    if font is None:
        font = ImageFont.truetype(path, size=size_px, index=face_index)
        _font_object_cache[cache_key] = font
    style = _norm_family(font.getname()[1])
    fake_bold = bold and "bold" not in style
    return font, fake_bold


# ---------------------------------------------------------------------------
# 値の書式
# ---------------------------------------------------------------------------

def format_cell_value(value: Any, number_format: str | None) -> str:
    """セル値を表示文字列にする（カルテで使う書式のみ対応）。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)):
        fmt = (number_format or "General").strip()
        if "%" in fmt:
            digits = 2 if "0.00" in fmt else 0
            return f"{value * 100:.{digits}f}%"
        if "#,##0" in fmt:
            return f"{value:,.0f}"
        if fmt == "0":
            return f"{value:.0f}"
        if float(value).is_integer():
            return str(int(value))
        return str(value)
    return str(value)


# ---------------------------------------------------------------------------
# 折り返し
# ---------------------------------------------------------------------------

def _tokenize_for_wrap(line: str) -> list[str]:
    """折り返し単位に分割する。英単語は塊、CJKは1文字ずつ。"""
    return re.findall(r"[A-Za-z0-9,.\-+%#()'\"]+\s*|\s+|.", line)


def wrap_line(line: str, font: Any, max_width: int, draw: ImageDraw.ImageDraw) -> list[str]:
    if not line:
        return [""]
    if draw.textlength(line, font=font) <= max_width:
        return [line]
    wrapped: list[str] = []
    current = ""
    for token in _tokenize_for_wrap(line):
        candidate = current + token
        if current and draw.textlength(candidate, font=font) > max_width:
            wrapped.append(current.rstrip())
            current = token.lstrip()
        else:
            current = candidate
    if current:
        wrapped.append(current.rstrip())
    return wrapped or [""]


# ---------------------------------------------------------------------------
# レンダラー本体
# ---------------------------------------------------------------------------

class _SheetGeometry:
    """列・行のピクセル境界（レンジ先頭からの累積座標）。"""

    def __init__(self, ws: Any, min_col: int, min_row: int, max_col: int, max_row: int, scale: int) -> None:
        default_width = ws.sheet_format.defaultColWidth or DEFAULT_COL_WIDTH
        default_height = ws.sheet_format.defaultRowHeight or DEFAULT_ROW_HEIGHT_PT

        self.col_x: dict[int, int] = {}
        x = 0
        for col in range(min_col, max_col + 2):
            self.col_x[col] = x
            if col > max_col:
                break
            dim = ws.column_dimensions.get(_col_letter(col))
            if dim is not None and dim.hidden:
                width_px = 0
            elif dim is not None and dim.width is not None:
                width_px = int(round(dim.width * CHAR_PIXELS)) + COL_PADDING_PIXELS
            else:
                width_px = int(round(default_width * CHAR_PIXELS)) + COL_PADDING_PIXELS
            x += width_px * scale

        self.row_y: dict[int, int] = {}
        y = 0
        for row in range(min_row, max_row + 2):
            self.row_y[row] = y
            if row > max_row:
                break
            dim = ws.row_dimensions.get(row)
            if dim is not None and dim.hidden:
                height_px = 0
            elif dim is not None and dim.height is not None:
                height_px = int(round(dim.height * POINTS_TO_PIXELS))
            else:
                height_px = int(round(default_height * POINTS_TO_PIXELS))
            y += height_px * scale

        self.width = self.col_x[max_col + 1]
        self.height = self.row_y[max_row + 1]

    def cell_rect(self, col: int, row: int, col_span_end: int | None = None, row_span_end: int | None = None) -> tuple[int, int, int, int]:
        x1 = self.col_x[col]
        y1 = self.row_y[row]
        x2 = self.col_x[(col_span_end or col) + 1]
        y2 = self.row_y[(row_span_end or row) + 1]
        return x1, y1, x2, y2


def _col_letter(col: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(col)


def render_range_to_png(
    xlsx_path: str | Path,
    sheet_name: str,
    cell_range: str,
    png_path: str | Path,
    *,
    scale: int = 2,
) -> Path:
    """1シートの矩形範囲をPNGに描画して保存する。"""
    xlsx_path = Path(xlsx_path)
    png_path = Path(png_path)
    theme_colors = load_theme_colors(xlsx_path)

    workbook = load_workbook(xlsx_path, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"シートがありません: {sheet_name}")
        ws = workbook[sheet_name]
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        geometry = _SheetGeometry(ws, min_col, min_row, max_col, max_row, scale)

        image = Image.new("RGB", (max(geometry.width, 1), max(geometry.height, 1)), "white")
        draw = ImageDraw.Draw(image)

        # 結合セル: 左上セル→結合範囲、内部セル→スキップ対象。
        merged_root: dict[tuple[int, int], tuple[int, int, int, int]] = {}
        merged_covered: set[tuple[int, int]] = set()
        for merged in ws.merged_cells.ranges:
            bounds = merged.bounds  # (min_col, min_row, max_col, max_row)
            merged_root[(bounds[0], bounds[1])] = bounds
            for row in range(bounds[1], bounds[3] + 1):
                for col in range(bounds[0], bounds[2] + 1):
                    if (col, row) != (bounds[0], bounds[1]):
                        merged_covered.add((col, row))

        def target_rect(col: int, row: int) -> tuple[int, int, int, int]:
            bounds = merged_root.get((col, row))
            if bounds:
                return geometry.cell_rect(
                    col, row,
                    col_span_end=min(bounds[2], max_col),
                    row_span_end=min(bounds[3], max_row),
                )
            return geometry.cell_rect(col, row)

        # 1) 塗りつぶし
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                if (col, row) in merged_covered:
                    continue
                cell = ws.cell(row=row, column=col)
                fill = cell.fill
                if fill is None or fill.fill_type != "solid":
                    continue
                color = resolve_color(fill.fgColor, theme_colors)
                if color is None:
                    continue
                x1, y1, x2, y2 = target_rect(col, row)
                if x2 > x1 and y2 > y1:
                    draw.rectangle([x1, y1, x2 - 1, y2 - 1], fill=color)

        # 2) 罫線
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                cell = ws.cell(row=row, column=col)
                border = cell.border
                if border is None:
                    continue
                x1, y1, x2, y2 = geometry.cell_rect(col, row)
                for side, segment in (
                    (border.top, (x1, y1, x2, y1)),
                    (border.bottom, (x1, y2, x2, y2)),
                    (border.left, (x1, y1, x1, y2)),
                    (border.right, (x2, y1, x2, y2)),
                ):
                    if side is None or not side.style:
                        continue
                    width = BORDER_WIDTHS.get(side.style, 1) * scale
                    color = resolve_color(side.color, theme_colors) or (0, 0, 0)
                    draw.line(segment, fill=color, width=width)

        # 3) テキスト
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                if (col, row) in merged_covered:
                    continue
                cell = ws.cell(row=row, column=col)
                text = format_cell_value(cell.value, cell.number_format)
                if text == "":
                    continue
                _draw_cell_text(
                    image, draw, ws, geometry, theme_colors,
                    cell, text, target_rect(col, row),
                    col, row, min_col, max_col, merged_root, merged_covered, scale,
                )
    finally:
        workbook.close()

    png_path.parent.mkdir(parents=True, exist_ok=True)
    image.save(png_path)
    return png_path


def _draw_cell_text(
    image: Image.Image,
    draw: ImageDraw.ImageDraw,
    ws: Any,
    geometry: _SheetGeometry,
    theme_colors: dict[str, str],
    cell: Any,
    text: str,
    cell_rect: tuple[int, int, int, int],
    col: int,
    row: int,
    min_col: int,
    max_col: int,
    merged_root: dict,
    merged_covered: set,
    scale: int,
) -> None:
    font_style = cell.font
    size_pt = float(font_style.size or 11.0)
    size_px = max(int(round(size_pt * POINTS_TO_PIXELS * scale)), 1)
    font, fake_bold = get_font(str(font_style.name or "Calibri"), size_px, bool(font_style.bold), text)
    color = resolve_color(font_style.color, theme_colors) or (0, 0, 0)

    alignment = cell.alignment
    wrap = bool(alignment.wrap_text)
    horizontal = alignment.horizontal
    if horizontal in (None, "general"):
        horizontal = "right" if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool) else "left"
    vertical = alignment.vertical or "bottom"

    pad = 3 * scale
    x1, y1, x2, y2 = cell_rect
    inner_width = max(x2 - x1 - 2 * pad, 1)

    # 非折り返しテキストは、隣が空セルなら Excel 同様はみ出して表示する。
    allowed_x1, allowed_x2 = x1, x2
    if not wrap:
        span_end = merged_root.get((col, row), (col, row, col, row))[2]
        if horizontal in ("left", "center"):
            next_col = span_end + 1
            while next_col <= max_col and _cell_is_vacant(ws, next_col, row, merged_covered, merged_root):
                allowed_x2 = geometry.col_x[next_col + 1]
                next_col += 1
        if horizontal in ("right", "center"):
            prev_col = col - 1
            while prev_col >= min_col and _cell_is_vacant(ws, prev_col, row, merged_covered, merged_root):
                allowed_x1 = geometry.col_x[prev_col]
                prev_col -= 1

    lines: list[str] = []
    for raw_line in text.split("\n"):
        if wrap:
            lines.extend(wrap_line(raw_line, font, inner_width, draw))
        else:
            lines.append(raw_line)

    ascent, descent = font.getmetrics()
    line_height = int((ascent + descent) * 1.02)
    block_height = line_height * len(lines)

    if vertical == "top":
        y_start = y1 + pad
    elif vertical in ("center", "justify"):
        y_start = y1 + max((y2 - y1 - block_height) // 2, pad // 2)
    else:  # bottom
        y_start = y2 - pad // 2 - block_height

    # クリッピング領域に描いてから貼り付ける。
    clip_x1, clip_y1 = allowed_x1, y1
    clip_x2, clip_y2 = allowed_x2, y2
    clip_width = clip_x2 - clip_x1
    clip_height = clip_y2 - clip_y1
    if clip_width <= 0 or clip_height <= 0:
        return
    layer = Image.new("RGBA", (clip_width, clip_height), (0, 0, 0, 0))
    layer_draw = ImageDraw.Draw(layer)
    stroke = 1 if fake_bold else 0

    for index, line in enumerate(lines):
        line_width = layer_draw.textlength(line, font=font)
        if horizontal == "center":
            x_text = x1 + (x2 - x1 - line_width) / 2
        elif horizontal == "right":
            x_text = x2 - pad - line_width
        else:
            x_text = x1 + pad
        y_text = y_start + index * line_height
        if y_text + line_height < clip_y1 or y_text > clip_y2:
            continue
        layer_draw.text(
            (x_text - clip_x1, y_text - clip_y1),
            line,
            font=font,
            fill=color + (255,),
            stroke_width=stroke,
            stroke_fill=color + (255,),
        )
    image.paste(layer, (clip_x1, clip_y1), layer)


def _cell_is_vacant(ws: Any, col: int, row: int, merged_covered: set, merged_root: dict) -> bool:
    """はみ出し表示のために「空きセルか」を判定する。"""
    if (col, row) in merged_covered or (col, row) in merged_root:
        return False
    value = ws.cell(row=row, column=col).value
    return value is None or str(value) == ""
